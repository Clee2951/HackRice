from io import BytesIO
from pathlib import Path
from fastapi import HTTPException
from pypdf import PdfReader
from docx import Document as DocxDocument
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from backend.core.config import settings
from backend.ai import prompts
from backend.schemas.study import ObjectiveSet
from backend.models.document import Document, DocumentObject
from backend.services import storage


def extract_pages(filename, data):
    suffix = Path(filename).suffix.lower()
    if not data or len(data) > settings.MAX_UPLOAD_BYTES:
        raise HTTPException(413, "Upload must contain data and be at most 10 MB")
    try:
        if suffix == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise ValueError("Invalid PDF")
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise HTTPException(422, "Upload an unencrypted PDF")
            if len(reader.pages) > settings.MAX_PDF_PAGES:
                raise HTTPException(413, "Split the PDF into sections of at most 50 pages")
            pages = [{"page": i + 1, "text": p.extract_text() or ""} for i, p in enumerate(reader.pages)]
            # Reject scans rather than silently generating objectives from incomplete text.
            if any(len(p["text"].strip()) < 20 for p in pages):
                raise HTTPException(422, "Some pages have little extractable text. Use a text PDF or a transcript")
            media = "application/pdf"
        elif suffix in {".txt", ".md"}:
            pages = [{"page": 1, "text": data.decode("utf-8-sig")}]
            media = "text/plain"
        elif suffix == ".docx":
            doc = DocxDocument(BytesIO(data))
            full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            pages = [{"page": 1, "text": full_text}]
            media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            raise HTTPException(415, "Supported formats: .docx, text-based PDF, UTF-8 .txt, .md")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, "Cannot read this file. Export a text PDF or UTF-8 text file")
    chars = sum(len(p["text"]) for p in pages)
    if chars < 40:
        raise HTTPException(422, "Not enough text to study")
    if chars > settings.MAX_DOCUMENT_CHARS:
        raise HTTPException(413, "Split the document into sections under 60,000 characters")
    return pages, media


def check_pages(numbers, pages):
    valid = {p["page"] for p in pages}
    if not set(numbers).issubset(valid):
        raise HTTPException(502, "AI cited an unknown page; please retry")


def build_objectives(pages, ai):
    result = ai.structured(prompts.EXTRACT, {"source": pages}, ObjectiveSet)
    objectives = []
    for index, obj in enumerate(result.objectives, 1):
        check_pages(obj.source_pages, pages)
        objectives.append({"id": f"c{index}", **obj.model_dump()})
    return objectives


def content_for_pages(pages, source_pages):
    # The grounding text stored on each document_objects row: the raw text
    # of just the pages that objective cites, not the whole document.
    wanted = set(source_pages)
    return "\n\n".join(p["text"] for p in pages if p["page"] in wanted)


def next_document_id(db, uid):
    # documents.document_id is scoped PER USER, not global -- there's no
    # native "auto-increment scoped to another column" in MySQL/SQLite, so
    # compute max+1 ourselves. create_document() retries on the rare race
    # (two uploads from the same account at the same instant).
    current = db.query(func.max(Document.document_id)).filter(Document.uid == uid).scalar()
    return (current or 0) + 1


def create_document(db, uid, title, media_type, data, pages, objectives):
    # One transaction: allocate the next per-user document_id, insert the
    # document row, then one document_objects row per AI-extracted
    # objective (object_id = its 1-based position -- see build_objectives,
    # which already numbers them "c1", "c2", ... in the same order; a
    # brand-new document has no existing rows to collide with, so this
    # doesn't need its own next_id lookup the way document_id does).
    use_object_storage = storage.configured()
    for _ in range(3):
        document_id = next_document_id(db, uid)
        doc = Document(uid=uid, document_id=document_id, document_title=title,
                       media_type=media_type, progress={})
        db.add(doc)
        try:
            db.flush()
        except IntegrityError:
            db.rollback()
            continue
        # Only write to Vultr once the id is actually ours -- the flush
        # above is what settles the race with a concurrent upload, and
        # uploading before it would orphan an object on every retry.
        if use_object_storage:
            doc.storage_key = storage.put_bytes(uid, document_id, title, data, media_type)
        else:
            doc.original = data
        for index, obj in enumerate(objectives, 1):
            db.add(DocumentObject(uid=uid, document_id=document_id, object_id=index,
                                  object_title=obj["title"], expected_points=obj["expected_points"],
                                  source_page=obj["source_pages"],
                                  content=content_for_pages(pages, obj["source_pages"])))
        try:
            db.commit()
        except Exception:
            # The object is already in the bucket at this point; without
            # this cleanup a failed commit leaves a file no row refers to.
            db.rollback()
            if doc.storage_key:
                storage.delete(doc.storage_key)
            raise
        return doc
    raise HTTPException(409, "Could not allocate a document slot, please retry")


def object_view(obj):
    # Reconstructs the "cN" id scheme the rest of the app (progress dict
    # keys, objective_ids, AssessmentItem.concept_id) already expects, so
    # only storage changed -- not the API shape or the AI-facing contract.
    # `content` is extra (not part of the old objectives blob) -- used as
    # AI grounding context in place of the old whole-document `pages`.
    return {"id": f"c{obj.object_id}", "title": obj.object_title,
            "expected_points": obj.expected_points, "source_pages": obj.source_page,
            "content": obj.content}


def valid_cited_pages(objects):
    return {page for obj in objects for page in (obj.source_page or [])}


def check_cited_pages(numbers, objects):
    if not set(numbers).issubset(valid_cited_pages(objects)):
        raise HTTPException(502, "AI cited an unknown page; please retry")
